#!/usr/bin/env python3
"""Check a protocol review's working files and build its tables.

Usage:
    python review_tables.py --ledger <ledger.yaml> --map <protocol_map.yaml>
        --judgements <judgements.yaml> --prefix <Protocol_Name> --outdir <workspace>
        [--register <clinical-evidence-register.csv> --model-id "..." --skill-version "..."]

Inputs (written by the protocol-reviewer skill in <workspace>/.clinical-evidence/):

  protocol_map.yaml
    protocol: {title, version, date, institution, domain, population}
    questions:  - {id: Q1, text: "..."}
    statements: - {id: S1, section: "3.2 Rituximab", text: "...", kind: dose, questions: [Q1]}

  judgements.yaml  (list)
    - id: J1
      statement: S1            # null only for verdict new_addition
      question: Q1
      verdict: major_update    # aligned | minor_update | major_update | new_addition | remove
      recommendation: "..."
      evidence: [4, 7]         # ledger ref_ids
      grade: "BTS Grade 1C"    # optional, grade.display verbatim
      safety: false            # patient-safety relevant change
      commissioning: false     # needs commissioner approval / business case
      confidence: high         # high | moderate | low
      second_review: {status: agree, note: "", resolution: ""}   # after second review

Checks (any ERROR → exit 1, nothing written): unique ids; statements and questions
referenced exist; every statement has at least one judgement; verdict and confidence
values valid; every evidence ref_id exists in the ledger and was not excluded by
verification; a second-review disagreement has a written resolution.

Outputs:
  <prefix>_Evidence_Table.csv        one row per cited source (R-friendly)
  <prefix>_Traceability.csv          statement → question → evidence → verdict
  <prefix>_Evidence_Table.xlsx       both tables as sheets (only if openpyxl is installed)
  .clinical-evidence/traceability.md Markdown matrix for the review's appendix
  register row                       appended when --register is given

Exit codes: 0 — ok; 1 — consistency errors; 2 — a file is missing or unreadable.
"""

from __future__ import annotations

import argparse
import csv
import json
import sys
from datetime import date
from pathlib import Path
from typing import Any

try:
    import yaml
except ImportError:
    sys.stderr.write("ERROR: PyYAML is not installed. Install with: pip install pyyaml\n")
    sys.exit(2)

VERDICTS = {
    "aligned": "Aligned",
    "minor_update": "Minor update",
    "major_update": "Major update",
    "new_addition": "New addition",
    "remove": "Remove",
}
CONFIDENCE = {"high", "moderate", "low"}
SECOND_REVIEW = {"agree", "disagree", "unsupported"}

EVIDENCE_COLUMNS = [
    "ref_id", "type", "citation", "title", "journal", "year", "pmid", "doi", "study_design",
    "population", "sample_size", "certainty", "key_finding", "questions", "cited_in",
    "verification",
]
TRACE_COLUMNS = [
    "judgement", "section", "statement_id", "statement", "kind", "question_id", "question",
    "verdict", "recommendation", "evidence", "evidence_citations", "grade", "safety",
    "commissioning", "confidence", "second_review", "second_review_note", "resolution",
]
REGISTER_COLUMNS = [
    "review_date", "protocol_name", "protocol_version", "clinical_domain", "skill_version",
    "model_id", "total_references", "guidelines_consulted", "recommendations_aligned",
    "recommendations_minor_update", "recommendations_major_update",
    "recommendations_new_addition", "recommendations_remove", "references_excluded",
    "second_review_disagreements", "mdt_outcome", "appraiser", "notes",
]


def _load(path: Path) -> Any:
    with path.open("r", encoding="utf-8") as f:
        return yaml.safe_load(f)


def _s(v: Any) -> str:
    return "" if v is None else str(v).strip()


def _join(items: Any) -> str:
    return "; ".join(_s(i) for i in (items or []))


def _citation(kind: str, entry: dict) -> str:
    if kind == "guideline":
        return f"{_s(entry.get('organisation'))} {_s(entry.get('year'))}".strip()
    author = _s(entry.get("first_author")) or _s(entry.get("authors")).split(",")[0]
    return f"{author} {_s(entry.get('year'))}".strip()


def index_ledger(ledger: dict) -> dict[int, tuple[str, dict]]:
    index: dict[int, tuple[str, dict]] = {}
    for kind, section in (("guideline", "guidelines"), ("article", "references"), ("preprint", "preprints")):
        for e in ledger.get(section) or []:
            if isinstance(e, dict) and isinstance(e.get("ref_id"), int):
                index[e["ref_id"]] = (kind, e)
    return index


def check(ledger: dict, pmap: dict, judgements: list) -> tuple[list[str], list[str]]:
    errors: list[str] = []
    warnings: list[str] = []
    index = index_ledger(ledger)
    excluded = {e.get("ref_id") for e in ledger.get("excluded_references") or [] if isinstance(e, dict)}

    questions = {q.get("id"): q for q in pmap.get("questions") or [] if isinstance(q, dict)}
    if len(questions) != len(pmap.get("questions") or []):
        errors.append("protocol_map: question ids must be present and unique")
    statements = {s.get("id"): s for s in pmap.get("statements") or [] if isinstance(s, dict)}
    if len(statements) != len(pmap.get("statements") or []):
        errors.append("protocol_map: statement ids must be present and unique")
    for sid, s in statements.items():
        for q in s.get("questions") or []:
            if q not in questions:
                errors.append(f"statement {sid}: unknown question {q}")

    seen, covered = set(), set()
    for i, j in enumerate(judgements):
        jid = j.get("id") or f"judgements[{i}]"
        if jid in seen:
            errors.append(f"{jid}: duplicate judgement id")
        seen.add(jid)
        verdict = j.get("verdict")
        if verdict not in VERDICTS:
            errors.append(f"{jid}: verdict must be one of {sorted(VERDICTS)}, got {verdict!r}")
        sid = j.get("statement")
        if sid is None:
            if verdict != "new_addition":
                errors.append(f"{jid}: only new_addition judgements may have no statement")
        elif sid not in statements:
            errors.append(f"{jid}: unknown statement {sid}")
        else:
            covered.add(sid)
        if j.get("question") is not None and j["question"] not in questions:
            errors.append(f"{jid}: unknown question {j['question']}")
        if not _s(j.get("recommendation")):
            errors.append(f"{jid}: recommendation is empty")
        if j.get("confidence") not in CONFIDENCE:
            errors.append(f"{jid}: confidence must be one of {sorted(CONFIDENCE)}")
        evidence = j.get("evidence") or []
        if not evidence and verdict != "aligned":
            warnings.append(f"{jid}: {VERDICTS.get(verdict, verdict)} with no cited evidence")
        for ref_id in evidence:
            if ref_id in excluded:
                errors.append(f"{jid}: cites ref_id {ref_id}, which was excluded by verification")
            elif ref_id not in index:
                errors.append(f"{jid}: cites ref_id {ref_id}, which is not in the ledger")
        sr = j.get("second_review")
        practice_changing = verdict in {"major_update", "new_addition", "remove"} or j.get("safety")
        if not sr:
            if practice_changing:
                warnings.append(f"{jid}: practice-changing but not second-reviewed")
        elif sr.get("status") not in SECOND_REVIEW:
            errors.append(f"{jid}: second_review.status must be one of {sorted(SECOND_REVIEW)}")
        elif sr["status"] != "agree" and not _s(sr.get("resolution")):
            errors.append(f"{jid}: second reviewer {sr['status']} — a resolution is required")

    for sid in statements:
        if sid not in covered:
            errors.append(f"statement {sid} has no judgement")
    return errors, warnings


def build(ledger: dict, pmap: dict, judgements: list) -> tuple[list[dict], list[dict], dict]:
    index = index_ledger(ledger)
    questions = {q["id"]: q for q in pmap.get("questions") or []}
    statements = {s["id"]: s for s in pmap.get("statements") or []}

    cited_in: dict[int, list[str]] = {}
    trace = []
    for j in judgements:
        for ref_id in j.get("evidence") or []:
            cited_in.setdefault(ref_id, []).append(j["id"])
        st = statements.get(j.get("statement")) or {}
        q = questions.get(j.get("question")) or {}
        sr = j.get("second_review") or {}
        trace.append({
            "judgement": j["id"], "section": _s(st.get("section") or j.get("section")),
            "statement_id": _s(j.get("statement")), "statement": _s(st.get("text")),
            "kind": _s(st.get("kind")), "question_id": _s(j.get("question")),
            "question": _s(q.get("text")), "verdict": VERDICTS[j["verdict"]],
            "recommendation": _s(j.get("recommendation")),
            "evidence": _join(j.get("evidence")),
            "evidence_citations": _join(_citation(*index[r]) for r in j.get("evidence") or []),
            "grade": _s(j.get("grade")), "safety": "yes" if j.get("safety") else "no",
            "commissioning": "yes" if j.get("commissioning") else "no",
            "confidence": _s(j.get("confidence")), "second_review": _s(sr.get("status")),
            "second_review_note": _s(sr.get("note")), "resolution": _s(sr.get("resolution")),
        })

    evidence = []
    for ref_id in sorted(cited_in):
        kind, e = index[ref_id]
        if kind == "guideline":
            finding = " | ".join(
                f"{_s(r.get('text'))} ({_s((r.get('grade') or {}).get('display'))})"
                for r in e.get("key_recommendations") or [])
        else:
            finding = _s(e.get("key_finding"))
        evidence.append({
            "ref_id": ref_id, "type": kind, "citation": _citation(kind, e),
            "title": _s(e.get("title")), "journal": _s(e.get("journal") or e.get("server")),
            "year": _s(e.get("year")), "pmid": _s(e.get("pmid")), "doi": _s(e.get("doi")),
            "study_design": _s(e.get("study_design")), "population": _s(e.get("population")),
            "sample_size": _s(e.get("sample_size")), "certainty": _s(e.get("certainty")),
            "key_finding": finding, "questions": _join(e.get("questions")),
            "cited_in": _join(cited_in[ref_id]),
            "verification": _s((e.get("integrity") or {}).get("status")) or ("n/a" if kind == "guideline" else ""),
        })

    counts = {v: sum(1 for j in judgements if j["verdict"] == v) for v in VERDICTS}
    counts["references_cited"] = len(cited_in)
    counts["references_excluded"] = len(ledger.get("excluded_references") or [])
    counts["second_review_disagreements"] = sum(
        1 for j in judgements if (j.get("second_review") or {}).get("status") in {"disagree", "unsupported"})
    counts["safety_flags"] = sum(1 for j in judgements if j.get("safety"))
    return evidence, trace, counts


def _write_csv(path: Path, columns: list[str], rows: list[dict]) -> None:
    with path.open("w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=columns)
        writer.writeheader()
        writer.writerows(rows)


def _write_xlsx(path: Path, sheets: list[tuple[str, list[str], list[dict]]]) -> bool:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font
    except ImportError:
        return False
    wb = Workbook()
    wb.remove(wb.active)
    for title, columns, rows in sheets:
        ws = wb.create_sheet(title)
        ws.append(columns)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for row in rows:
            ws.append([row.get(c, "") for c in columns])
        ws.freeze_panes = "A2"
        for col in ws.columns:
            width = min(60, max(10, *(len(str(c.value or "")) for c in col)))
            ws.column_dimensions[col[0].column_letter].width = width
            for c in col[1:]:
                c.alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(path)
    return True


def _markdown(trace: list[dict]) -> str:
    cols = ["judgement", "section", "statement", "question_id", "verdict", "evidence_citations",
            "grade", "second_review"]
    head = ["#", "Section", "Protocol statement", "Q", "Assessment", "Evidence", "Grade", "2nd review"]
    lines = ["| " + " | ".join(head) + " |", "|" + "---|" * len(head)]
    for r in trace:
        lines.append("| " + " | ".join(_s(r[c]).replace("|", "/").replace("\n", " ") for c in cols) + " |")
    return "\n".join(lines) + "\n"


def append_register(path: Path, row: dict) -> None:
    existing_rows, header = [], list(REGISTER_COLUMNS)
    if path.exists():
        with path.open("r", encoding="utf-8", newline="") as f:
            reader = csv.DictReader(f)
            old = reader.fieldnames or []
            existing_rows = list(reader)
        header = old + [c for c in REGISTER_COLUMNS if c not in old]
    with path.open("w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=header, restval="")
        writer.writeheader()
        writer.writerows(existing_rows)
        writer.writerow(row)


def main() -> int:
    p = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    p.add_argument("--ledger", type=Path, required=True)
    p.add_argument("--map", type=Path, required=True, dest="pmap")
    p.add_argument("--judgements", type=Path, required=True)
    p.add_argument("--prefix", required=True)
    p.add_argument("--outdir", type=Path, required=True)
    p.add_argument("--register", type=Path)
    p.add_argument("--model-id", default="")
    p.add_argument("--skill-version", default="")
    args = p.parse_args()

    try:
        ledger, pmap, judgements = _load(args.ledger), _load(args.pmap), _load(args.judgements)
    except (OSError, yaml.YAMLError) as e:
        sys.stderr.write(f"ERROR: {e}\n")
        return 2
    if isinstance(judgements, dict):
        judgements = judgements.get("judgements")
    if not isinstance(ledger, dict) or not isinstance(pmap, dict) or not isinstance(judgements, list):
        sys.stderr.write("ERROR: ledger and protocol map must be mappings; judgements a list\n")
        return 2

    errors, warnings = check(ledger, pmap, judgements)
    for w in warnings:
        print(f"WARN: {w}")
    if errors:
        for e in errors:
            print(f"ERROR: {e}")
        print(f"\n{len(errors)} error(s) — no tables written.")
        return 1

    evidence, trace, counts = build(ledger, pmap, judgements)
    args.outdir.mkdir(parents=True, exist_ok=True)
    ev_csv = args.outdir / f"{args.prefix}_Evidence_Table.csv"
    tr_csv = args.outdir / f"{args.prefix}_Traceability.csv"
    _write_csv(ev_csv, EVIDENCE_COLUMNS, evidence)
    _write_csv(tr_csv, TRACE_COLUMNS, trace)
    xlsx = args.outdir / f"{args.prefix}_Evidence_Table.xlsx"
    has_xlsx = _write_xlsx(xlsx, [("Evidence", EVIDENCE_COLUMNS, evidence),
                                  ("Traceability", TRACE_COLUMNS, trace)])
    md_path = args.outdir / ".clinical-evidence" / "traceability.md"
    md_path.parent.mkdir(parents=True, exist_ok=True)
    md_path.write_text(_markdown(trace), encoding="utf-8")

    if args.register:
        proto = pmap.get("protocol") or {}
        append_register(args.register, {
            "review_date": date.today().isoformat(), "protocol_name": _s(proto.get("title")),
            "protocol_version": _s(proto.get("version")), "clinical_domain": _s(proto.get("domain")),
            "skill_version": args.skill_version, "model_id": args.model_id,
            "total_references": counts["references_cited"],
            "guidelines_consulted": "; ".join(
                f"{_s(g.get('organisation'))} {_s(g.get('year'))}".strip()
                for g in ledger.get("guidelines") or []),
            **{f"recommendations_{v}": counts[v] for v in VERDICTS},
            "references_excluded": counts["references_excluded"],
            "second_review_disagreements": counts["second_review_disagreements"],
        })

    print(json.dumps(counts))
    print(f"Wrote {ev_csv.name}, {tr_csv.name}"
          + (f", {xlsx.name}" if has_xlsx else " (install openpyxl for the .xlsx)")
          + f", {md_path}" + (f"; register row appended to {args.register}" if args.register else ""))
    return 0


if __name__ == "__main__":
    sys.exit(main())
